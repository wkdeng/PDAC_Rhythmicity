"""Read Kronos time-course data and fit a fixed-period single-harmonic cosinor.

The Kronos workbook contains measured signals with no stated units.  This module
interprets its time values as hours and treats every finite row in a trace as an
observation.  It does not smooth, detrend, transform, or infer biological
replicates.  Consequently, its p-values are nominal OLS tests: they assume
independent, identically distributed Gaussian residuals and should not be used
as replicate-level biological inference.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats


_TIME_HEADERS = {"timepoint", "timepoints"}
_GENE_NAMES = {"bmal": "BMAL1", "bmal1": "BMAL1", "per2": "PER2"}


def _clean_header(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def _as_finite_number(value: Any, *, sheet: str, excel_row: int, column: int,
                      field: str) -> float | None:
    """Convert a populated numeric workbook cell or raise a useful error."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, (int, float, np.number)):
        raise ValueError(
            f"Unexpected nonnumeric {field} in {sheet}!R{excel_row}C{column}: {value!r}"
        )
    number = float(value)
    if not np.isfinite(number):
        raise ValueError(f"Non-finite {field} in {sheet}!R{excel_row}C{column}: {value!r}")
    return number


def _find_header_row(rows: list[tuple[Any, ...]], sheet: str) -> int:
    for index, row in enumerate(rows):
        if any(_clean_header(value) in _TIME_HEADERS for value in row):
            return index
    raise ValueError(f"No timepoint header found in worksheet {sheet!r}")


def read_kronos(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Read valid Kronos time/signal pairs and per-trace missingness QC.

    Valid observations have finite numeric time and signal values.  The returned
    observations have columns ``cell_line, gene, time_h, signal, excel_row,
    sheet``.  A populated nonnumeric cell and duplicate valid times within a
    cell-line/gene trace are errors rather than silently coerced values.
    """
    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    observations: list[dict[str, Any]] = []
    qc_rows: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        header_index = _find_header_row(rows, worksheet.title)
        header = rows[header_index]
        time_columns = [i for i, value in enumerate(header) if _clean_header(value) in _TIME_HEADERS]
        if not time_columns:
            raise ValueError(f"No timepoint columns found in worksheet {worksheet.title!r}")

        # A time column supplies all intervening signal columns, up to the next
        # time column.  This covers the shared-time and paired-time workbook layouts.
        trace_specs: list[tuple[int, int, str]] = []
        for time_position, time_column in enumerate(time_columns):
            next_time = time_columns[time_position + 1] if time_position + 1 < len(time_columns) else len(header)
            for signal_column in range(time_column + 1, next_time):
                gene_header = _clean_header(header[signal_column])
                if gene_header not in _GENE_NAMES:
                    raise ValueError(
                        f"Unexpected signal header in {worksheet.title}!R{header_index + 1}C{signal_column + 1}: "
                        f"{header[signal_column]!r}"
                    )
                trace_specs.append((time_column, signal_column, _GENE_NAMES[gene_header]))

        if not trace_specs:
            raise ValueError(f"No supported signal columns found in worksheet {worksheet.title!r}")

        for time_column, signal_column, gene in trace_specs:
            counts = {"n_sheet_rows": 0, "n_valid": 0, "n_missing_time": 0,
                      "n_missing_signal": 0, "n_both_missing": 0}
            trace_records: list[dict[str, Any]] = []
            for row_index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
                time_value = row[time_column] if time_column < len(row) else None
                signal_value = row[signal_column] if signal_column < len(row) else None
                time_h = _as_finite_number(time_value, sheet=worksheet.title, excel_row=row_index,
                                           column=time_column + 1, field="time")
                signal = _as_finite_number(signal_value, sheet=worksheet.title, excel_row=row_index,
                                           column=signal_column + 1, field="signal")
                counts["n_sheet_rows"] += 1
                if time_h is None and signal is None:
                    counts["n_both_missing"] += 1
                elif time_h is None:
                    counts["n_missing_time"] += 1
                elif signal is None:
                    counts["n_missing_signal"] += 1
                else:
                    counts["n_valid"] += 1
                    trace_records.append({"cell_line": worksheet.title, "gene": gene,
                                          "time_h": time_h, "signal": signal,
                                          "excel_row": row_index, "sheet": worksheet.title})
            trace_data = pd.DataFrame(trace_records)
            if not trace_data.empty and trace_data["time_h"].duplicated().any():
                repeated = trace_data.loc[trace_data["time_h"].duplicated(keep=False), "time_h"].tolist()
                raise ValueError(f"Duplicate valid times for {worksheet.title} {gene}: {repeated}")
            observations.extend(trace_records)
            qc_rows.append({"cell_line": worksheet.title, "gene": gene, "sheet": worksheet.title,
                            "time_column": time_column + 1, "signal_column": signal_column + 1,
                            **counts})

    columns = ["cell_line", "gene", "time_h", "signal", "excel_row", "sheet"]
    workbook.close()
    return pd.DataFrame(observations, columns=columns), pd.DataFrame(qc_rows)


def fit_single_trace(time: Any, signal: Any, period_h: float = 24) -> dict[str, float]:
    """Fit ``M + bc*cos(2πt/T) + bs*sin(2πt/T)`` to finite supplied values."""
    if not np.isfinite(period_h) or period_h <= 0:
        raise ValueError("period_h must be a positive finite number")
    time_array = np.asarray(time, dtype=float).reshape(-1)
    signal_array = np.asarray(signal, dtype=float).reshape(-1)
    if time_array.size != signal_array.size:
        raise ValueError("time and signal must have equal lengths")
    valid = np.isfinite(time_array) & np.isfinite(signal_array)
    time_array, signal_array = time_array[valid], signal_array[valid]
    n = int(signal_array.size)
    if n < 3:
        raise ValueError("At least three finite time/signal pairs are required for cosinor fitting")
    if np.unique(time_array).size != n:
        raise ValueError("Duplicate finite timepoints are not allowed within a cosinor trace")
    order = np.argsort(time_array)
    time_array, signal_array = time_array[order], signal_array[order]
    omega = 2 * np.pi / period_h
    design = np.column_stack((np.ones(n), np.cos(omega * time_array), np.sin(omega * time_array)))
    coefficients, _, rank, _ = np.linalg.lstsq(design, signal_array, rcond=None)
    if rank < 3:
        raise ValueError("Cosinor design is rank-deficient; supply timepoints spanning the harmonic")
    fitted = design @ coefficients
    residual = signal_array - fitted
    sse = float(np.sum(residual ** 2))
    sst = float(np.sum((signal_array - signal_array.mean()) ** 2))
    r_squared = float(1 - sse / sst) if sst > 0 else np.nan
    rmse = float(np.sqrt(sse / n))
    intercept, beta_cos, beta_sin = map(float, coefficients)
    amplitude = float(np.hypot(beta_cos, beta_sin))
    peak_phase_h = float((np.arctan2(beta_sin, beta_cos) * period_h / (2 * np.pi)) % period_h)
    df_residual = n - 3
    sse_null = float(np.sum((signal_array - signal_array.mean()) ** 2))
    near_constant = sse_null <= np.finfo(float).eps * max(float(np.sum(signal_array ** 2)), float(n))
    if near_constant:
        # A flat trace has no identified phase and cannot support a harmonic test.
        peak_phase_h = np.nan
        f_statistic = p_value = np.nan
    elif df_residual <= 0:
        f_statistic = p_value = np.nan
    elif sse <= np.finfo(float).eps * max(sse_null, 1.0):
        f_statistic, p_value = np.inf, 0.0
    else:
        f_statistic = max(0.0, ((sse_null - sse) / 2) / (sse / df_residual))
        p_value = float(stats.f.sf(f_statistic, 2, df_residual))
    residual_lag1 = (float(np.corrcoef(residual[:-1], residual[1:])[0, 1])
                     if n >= 3 and np.std(residual[:-1]) > 0 and np.std(residual[1:]) > 0 else np.nan)
    return {"n_observations": n, "period_h": float(period_h), "intercept": intercept,
            "beta_cos": beta_cos, "beta_sin": beta_sin, "amplitude": amplitude,
            "peak_phase_h": peak_phase_h, "acrophase_h": peak_phase_h,
            "mesor": intercept, "r_squared": r_squared, "rmse": rmse,
            "f_statistic": f_statistic, "p_value_nominal": p_value,
            "df_numerator": 2, "df_residual": df_residual,
            "residual_lag1_correlation": residual_lag1}


def _bh_adjust(p_values: pd.Series) -> pd.Series:
    valid = p_values.notna()
    adjusted = pd.Series(np.nan, index=p_values.index, dtype=float)
    if valid.any():
        values = p_values.loc[valid].to_numpy(float)
        order = np.argsort(values)
        ranked = values[order] * len(values) / np.arange(1, len(values) + 1)
        ranked = np.minimum.accumulate(ranked[::-1])[::-1]
        result = np.empty_like(ranked)
        result[order] = np.minimum(ranked, 1.0)
        adjusted.loc[valid] = result
    return adjusted


def fit_cosinor(data: pd.DataFrame, period_h: float = 24) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Fit each cell-line/gene trace and return summary, observed rows, and dense curves.

    ``data`` must contain the six columns emitted by :func:`read_kronos`.  The
    observed output appends raw and ddof=0 z-score fitted values/residuals.  The
    curve output provides 1,000 predictions per trace across that trace's full
    observed recording range, both in raw signal units and using each trace's
    ddof=0 signal z-score scale.
    """
    required = {"cell_line", "gene", "time_h", "signal", "excel_row", "sheet"}
    missing = required.difference(data.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")
    summary_rows: list[dict[str, Any]] = []
    observed_parts: list[pd.DataFrame] = []
    curve_rows: list[dict[str, Any]] = []
    for (cell_line, gene), trace in data.groupby(["cell_line", "gene"], sort=True):
        trace = trace.copy().sort_values("time_h")
        result = fit_single_trace(trace["time_h"], trace["signal"], period_h=period_h)
        omega = 2 * np.pi / period_h
        fitted = (result["intercept"] + result["beta_cos"] * np.cos(omega * trace["time_h"].to_numpy()) +
                  result["beta_sin"] * np.sin(omega * trace["time_h"].to_numpy()))
        trace["fitted"] = fitted
        trace["residual"] = trace["signal"].to_numpy() - fitted
        signal_mean = float(trace["signal"].mean())
        signal_sd = float(trace["signal"].std(ddof=0))
        trace["signal_zscore"] = ((trace["signal"] - signal_mean) / signal_sd
                                  if signal_sd > 0 else np.nan)
        trace["fitted_zscore"] = ((trace["fitted"] - signal_mean) / signal_sd
                                  if signal_sd > 0 else np.nan)
        observed_parts.append(trace)
        summary_rows.append({"cell_line": cell_line, "gene": gene, "sheet": trace["sheet"].iloc[0],
                             "signal_mean": signal_mean, "signal_sd_ddof0": signal_sd, **result})
        grid = np.linspace(float(trace["time_h"].min()), float(trace["time_h"].max()), 1000)
        prediction = (result["intercept"] + result["beta_cos"] * np.cos(omega * grid) +
                      result["beta_sin"] * np.sin(omega * grid))
        for time_h, value in zip(grid, prediction):
            curve_rows.append({"cell_line": cell_line, "gene": gene, "sheet": trace["sheet"].iloc[0],
                               "time_h": float(time_h), "fitted": float(value),
                               "fitted_zscore": float((value - signal_mean) / signal_sd) if signal_sd > 0 else np.nan})
    summary = pd.DataFrame(summary_rows)
    summary["p_value_bh"] = _bh_adjust(summary["p_value_nominal"])
    observed = pd.concat(observed_parts, ignore_index=True) if observed_parts else data.copy()
    return summary, observed, pd.DataFrame(curve_rows)
